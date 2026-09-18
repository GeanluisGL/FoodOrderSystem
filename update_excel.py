import openpyxl
import sys
from datetime import datetime
import os

# 1. Configurar nombres de archivo y columnas
EXCEL_FILE = "historial_cambios.xlsx"
HEADERS = ["Fecha y Hora", "Usuario", "Rama", "Release", "Mensaje de Commit", "ID Commit"]

# 2. Capturar argumentos enviados por GitHub Actions
usuario = sys.argv[1]
rama = sys.argv[2].replace("refs/heads/", "").replace("refs/tags/", "")
release = sys.argv[3] if sys.argv[3] else "N/A"
commit_msg = sys.argv[4]
commit_id = sys.argv[5][:7] # Solo los primeros 7 caracteres
fecha_actual = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

# 3. Crear o abrir el archivo Excel
if not os.path.exists(EXCEL_FILE):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historial de Cambios"
    ws.append(HEADERS)
    # Dar formato negrita al encabezado
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
else:
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active

# 4. Insertar los nuevos datos en la última fila libre
ws.append([fecha_actual, usuario, rama, release, commit_msg, commit_id])

# 5. Guardar el archivo
wb.save(EXCEL_FILE)
print(f"Cambio registrado exitosamente para la rama: {rama}")
