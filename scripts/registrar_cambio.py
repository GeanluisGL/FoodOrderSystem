import os
import sys
from datetime import datetime

import openpyxl
from openpyxl.styles import Font


# ==========================================================
# 1. Configuración
# ==========================================================
EXCEL_FILE = "historial_cambios.xlsx"
SHEET_NAME = "Historial de Cambios"
HEADERS = [
    "Fecha y Hora",
    "Usuario",
    "Rama",
    "Release",
    "Mensaje de Commit",
    "ID Commit",
]
DATE_FORMAT = "YYYY-MM-DD HH:MM:SS"


# ==========================================================
# 2. Utilidades
# ==========================================================
def get_arg(index: int, default: str = "") -> str:
    """Devuelve sys.argv[index] o un valor por defecto si no existe o está vacío."""
    if len(sys.argv) > index:
        value = sys.argv[index]
        return value.strip() if value else default
    return default


def limpiar_ref(ref: str) -> str:
    """Convierte 'refs/heads/main' → 'main' y 'refs/tags/v1.0' → 'v1.0'."""
    return (
        ref.replace("refs/heads/", "")
           .replace("refs/tags/", "")
           .strip()
    )


def abrir_o_crear_workbook(ruta: str):
    """Abre el Excel si existe; si no, lo crea con encabezados en negrita."""
    if os.path.exists(ruta):
        wb = openpyxl.load_workbook(ruta)
        ws = wb.active
        # Asegurar que la hoja tenga el nombre correcto (por si acaso)
        ws.title = SHEET_NAME

        # Si por alguna razón no tiene encabezados, los agregamos
        if ws.max_row == 0 or ws.cell(row=1, column=1).value is None:
            ws.append(HEADERS)
            for cell in ws[1]:
                cell.font = Font(bold=True)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)

    return wb, ws


# ==========================================================
# 3. Captura de argumentos (enviados por GitHub Actions)
# ==========================================================
usuario    = get_arg(1, "desconocido")
rama       = limpiar_ref(get_arg(2, "desconocida"))
release    = get_arg(3, "N/A") or "N/A"
commit_msg = get_arg(4, "(sin mensaje)")
commit_id  = get_arg(5, "")[:7] or "N/A"
fecha_actual = datetime.now()


# ==========================================================
# 4. Abrir o crear el Excel
# ==========================================================
wb, ws = abrir_o_crear_workbook(EXCEL_FILE)


# ==========================================================
# 5. Insertar la nueva fila
# ==========================================================
ws.append([
    fecha_actual,   # se guarda como fecha real de Excel
    usuario,
    rama,
    release,
    commit_msg,
    commit_id,
])

# Dar formato de fecha a la celda recién insertada
nueva_fila = ws.max_row
ws.cell(row=nueva_fila, column=1).number_format = DATE_FORMAT

# Ajustar ancho de columnas (opcional pero agradable)
anchos = [20, 18, 15, 12, 50, 12]
for i, ancho in enumerate(anchos, start=1):
    letra = openpyxl.utils.get_column_letter(i)
    ws.column_dimensions[letra].width = ancho


# ==========================================================
# 6. Guardar y confirmar
# ==========================================================
try:
    wb.save(EXCEL_FILE)
    print(f"✅ Cambio registrado exitosamente para la rama: {rama}")
    print(f"   Usuario : {usuario}")
    print(f"   Commit  : {commit_id}")
    print(f"   Fecha   : {fecha_actual.strftime('%Y-%m-%d %H:%M:%S')}")
except PermissionError:
    print("❌ No se pudo guardar el archivo. ¿Está abierto en Excel?")
    sys.exit(1)
except Exception as e:
    print(f"❌ Error inesperado al guardar: {e}")
    sys.exit(1)
